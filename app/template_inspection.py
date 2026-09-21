from __future__ import annotations

from io import BytesIO
import re
from typing import Any

from docx import Document


KNOWN_FIELDS: list[tuple[str, str, tuple[str, ...], str, str]] = [
    ("complaints", "Жалобы", ("жалобы", "жалобы пациента"), "textarea", "patient"),
    ("anamnesis_morbi", "Anamnesis morbi", ("anamnesis morbi", "анамнез заболевания", "история настоящего заболевания"), "textarea", "both"),
    ("anamnesis_vitae", "Anamnesis vitae", ("anamnesis vitae", "анамнез жизни"), "textarea", "patient"),
    ("objective_status", "Объективный статус", ("объективный статус", "status praesens", "объективные данные"), "textarea", "doctor"),
    ("diagnosis_text", "Диагноз", ("диагноз", "предварительный диагноз", "клинический диагноз"), "text", "doctor"),
    ("examination_plan", "План обследования", ("план обследования", "план диагностики"), "textarea", "doctor"),
    ("treatment_plan", "План лечения", ("план лечения", "лечение", "план ведения"), "textarea", "doctor"),
    ("recommendations", "Рекомендации", ("рекомендации", "рекомендации врача"), "textarea", "doctor"),
    ("follow_up", "Повторный прием / наблюдение", ("повторный прием", "динамическое наблюдение", "наблюдение"), "textarea", "doctor"),
]


def _normalize(value: str) -> str:
    value = value.strip().casefold()
    value = re.sub(r"[\s\u00a0]+", " ", value)
    value = value.strip(" :;.-–—_")
    return value


def _collect_text(doc: Document) -> list[str]:
    values: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            values.append(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    values.append(text)
    return values


def inspect_docx_fields(raw: bytes) -> list[dict[str, Any]]:
    doc = Document(BytesIO(raw))
    texts = _collect_text(doc)
    normalized = [_normalize(item) for item in texts]

    out: list[dict[str, Any]] = []
    seen: set[str] = set()

    for field_id, label, aliases, field_type, ai_source in KNOWN_FIELDS:
        aliases_norm = tuple(_normalize(alias) for alias in aliases)
        matched = any(
            any(
                text == alias
                or text.startswith(alias + " ")
                or text.startswith(alias + ":")
                or alias in text[: max(len(alias) + 24, len(text))]
                for alias in aliases_norm
            )
            for text in normalized
        )
        if matched and field_id not in seen:
            out.append({
                "id": field_id,
                "label": label,
                "type": field_type,
                "ai_source": ai_source,
            })
            seen.add(field_id)

    if not out:
        out = [
            {"id": "complaints", "label": "Жалобы", "type": "textarea", "ai_source": "patient"},
            {"id": "objective_status", "label": "Объективный статус", "type": "textarea", "ai_source": "doctor"},
            {"id": "diagnosis_text", "label": "Диагноз", "type": "text", "ai_source": "doctor"},
            {"id": "recommendations", "label": "Рекомендации", "type": "textarea", "ai_source": "doctor"},
        ]

    return out
