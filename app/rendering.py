from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

import qrcode
from docx import Document
from docx.shared import Inches


TOKEN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


def replace_placeholders(text: str, values: dict[str, Any]) -> str:
    def repl(match: re.Match[str]) -> str:
        key = match.group(1)
        if key not in values:
            return match.group(0)
        value = values[key]
        return "" if value is None else str(value)
    return TOKEN.sub(repl, text)


def _replace_paragraph(paragraph, values: dict[str, Any]) -> None:
    full = "".join(run.text for run in paragraph.runs)
    replaced = replace_placeholders(full, values)
    if replaced == full:
        return
    if paragraph.runs:
        paragraph.runs[0].text = replaced
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = replaced


def render_docx(source: Path, destination: Path, values: dict[str, Any], qr_url: str | None = None) -> None:
    doc = Document(source)
    for paragraph in doc.paragraphs:
        _replace_paragraph(paragraph, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_paragraph(paragraph, values)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            _replace_paragraph(paragraph, values)
        for paragraph in section.footer.paragraphs:
            _replace_paragraph(paragraph, values)
        if qr_url:
            with tempfile.TemporaryDirectory() as td:
                qr_path = Path(td) / "qr.png"
                qrcode.make(qr_url).save(qr_path)
                p = section.footer.add_paragraph()
                p.add_run("Проверка документа: ")
                p.add_run().add_picture(str(qr_path), width=Inches(0.75))
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice is required for PDF rendering")
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td, str(docx_path)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        generated = Path(td) / f"{docx_path.stem}.pdf"
        if not generated.exists():
            raise RuntimeError("LibreOffice did not create PDF")
        shutil.copy2(generated, pdf_path)
