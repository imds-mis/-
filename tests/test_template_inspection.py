from io import BytesIO

from docx import Document

from app.template_inspection import inspect_docx_fields


def test_inspect_docx_detects_common_clinical_sections():
    doc = Document()
    doc.add_paragraph("Жалобы:")
    doc.add_paragraph("Anamnesis morbi:")
    doc.add_paragraph("Объективный статус:")
    doc.add_paragraph("Диагноз:")
    doc.add_paragraph("Рекомендации:")
    buf = BytesIO()
    doc.save(buf)

    fields = inspect_docx_fields(buf.getvalue())

    ids = [item["id"] for item in fields]
    assert "complaints" in ids
    assert "anamnesis_morbi" in ids
    assert "objective_status" in ids
    assert "diagnosis_text" in ids
    assert "recommendations" in ids


def test_inspect_docx_does_not_duplicate_same_section():
    doc = Document()
    doc.add_paragraph("Жалобы:")
    doc.add_paragraph("Жалобы пациента:")
    buf = BytesIO()
    doc.save(buf)

    fields = inspect_docx_fields(buf.getvalue())

    assert [item["id"] for item in fields].count("complaints") == 1
