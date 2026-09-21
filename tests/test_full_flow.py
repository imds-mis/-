import json
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.main import create_app


def make_docx(path: Path) -> bytes:
    doc = Document()
    doc.add_paragraph("Пациент: {{patient.full_name}}")
    doc.add_paragraph("Жалобы: {{complaints}}")
    doc.add_paragraph("Диагноз: {{diagnosis_text}}")
    doc.save(path)
    return path.read_bytes()


def test_full_medical_document_flow(tmp_path: Path):
    def fake_stt(audio: bytes, content_type: str) -> str:
        return "Боль внизу живота"

    app = create_app(
        database_url=f"sqlite:///{tmp_path / 'test.db'}",
        storage_root=tmp_path / "storage",
        qr_secret="super-secret",
        public_base_url="http://testserver",
        stt_url=None,
        stt_callable=fake_stt,
    )
    client = TestClient(app)
    headers = {
        "X-Tenant-ID": "t1", "X-User-ID": "u1", "X-Branch-ID": "b1",
        "X-Practitioner-ID": "p1", "X-Specialty-Code": "GYNE",
    }

    source = tmp_path / "source.docx"
    created = client.post("/v1/admin/templates", headers=headers,
        files={"file": ("source.docx", make_docx(source), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={
            "name": "Осмотр гинеколога",
            "fields_json": json.dumps([
                {"id": "complaints", "label": "Жалобы", "type": "textarea"},
                {"id": "diagnosis_text", "label": "Диагноз", "type": "text"},
            ]),
            "assignments_json": json.dumps([{"specialty_code": "GYNE"}]),
        }).json()["data"]
    assert client.post(f"/v1/admin/templates/{created['id']}/publish", headers=headers).status_code == 200

    eligible = client.get("/v1/doctor/templates", headers=headers).json()["data"]
    assert eligible[0]["id"] == created["id"]

    document = client.post("/v1/doctor/documents", headers=headers, json={
        "template_id": created["id"], "patient_id": "patient-1",
        "system_values": {"patient.full_name": "Иванова А.А."}
    }).json()["data"]

    speech = client.post(f"/v1/doctor/documents/{document['id']}/speech/complaints", headers=headers,
        files={"audio": ("dictation.webm", b"voice", "audio/webm")})
    assert speech.json()["data"]["field_id"] == "complaints"

    imported = client.post("/v1/admin/icd10/import", headers=headers,
        files={"file": ("icd.csv", "code,title\nN80,Эндометриоз\n".encode(), "text/csv")},
        data={"version": "10", "source": "authorized-test-source"})
    assert imported.status_code == 201
    assert client.get("/v1/doctor/icd10?q=эндомет", headers=headers).json()["data"][0]["code"] == "N80"

    protocol = client.post("/v1/admin/protocols", headers=headers, json={
        "title": "Эндометриоз: клинический протокол", "version": "1",
        "source_org": "МЗ РК", "status": "published", "icd_codes": ["N80"],
        "items": [{"type": "investigation", "title": "УЗИ органов малого таза"}]
    })
    assert protocol.status_code == 201
    assert client.get("/v1/doctor/protocols?icd_code=N80", headers=headers).json()["data"]

    client.patch(f"/v1/doctor/documents/{document['id']}/fields", headers=headers,
        json={"values": {"complaints": "Боль внизу живота", "diagnosis_text": "N80 — Эндометриоз"}})

    finalized = client.post(f"/v1/doctor/documents/{document['id']}/finalize", headers=headers)
    assert finalized.status_code == 200, finalized.text
    data = finalized.json()["data"]
    assert len(data["sha256"]) == 64

    second = client.post(f"/v1/doctor/documents/{document['id']}/finalize", headers=headers)
    assert second.json()["data"]["sha256"] == data["sha256"]

    pdf = client.get(f"/v1/doctor/documents/{document['id']}/pdf", headers=headers)
    assert pdf.content.startswith(b"%PDF")

    verified = client.get(data["verification_url"].replace("http://testserver", ""))
    assert verified.status_code == 200
    assert verified.content.startswith(b"%PDF")

    other = dict(headers)
    other["X-Tenant-ID"] = "t2"
    assert client.patch(f"/v1/doctor/documents/{document['id']}/fields", headers=other, json={"values": {"complaints": "x"}}).status_code == 404


def test_published_template_change_creates_new_version_and_old_document_keeps_snapshot(tmp_path: Path):
    app = create_app(database_url=f"sqlite:///{tmp_path/'db.sqlite'}", storage_root=tmp_path/"storage", qr_secret="s", public_base_url="http://testserver", stt_url=None)
    client = TestClient(app)
    h = {"X-Tenant-ID":"t1","X-User-ID":"u1","X-Branch-ID":"b1","X-Practitioner-ID":"p1","X-Specialty-Code":"GYNE"}
    source = tmp_path/"v1.docx"
    d = Document()
    d.add_paragraph("V1 {{complaints}}")
    d.save(source)
    t = client.post("/v1/admin/templates", headers=h,
        files={"file":("v1.docx",source.read_bytes(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"name":"T","fields_json":json.dumps([{"id":"complaints","type":"textarea"}]),"assignments_json":json.dumps([{"specialty_code":"GYNE"}])}
    ).json()["data"]["id"]
    client.post(f"/v1/admin/templates/{t}/publish", headers=h)
    first = client.post("/v1/doctor/documents", headers=h, json={"template_id":t,"patient_id":"p1"}).json()["data"]
    assert first["template_version"] == 1

    source2 = tmp_path/"v2.docx"
    d = Document()
    d.add_paragraph("V2 {{complaints}}")
    d.save(source2)
    created_v2 = client.post(f"/v1/admin/templates/{t}/versions", headers=h,
        files={"file":("v2.docx",source2.read_bytes(),"application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"fields_json":json.dumps([{"id":"complaints","type":"textarea"}])}
    )
    assert created_v2.status_code == 201, created_v2.text
    assert created_v2.json()["data"]["version"] == 2
    client.post(f"/v1/admin/templates/{t}/publish", headers=h)
    second = client.post("/v1/doctor/documents", headers=h, json={"template_id":t,"patient_id":"p2"}).json()["data"]
    assert second["template_version"] == 2
    assert first["template_version"] == 1


def test_doctor_profile_specialty_grants_template_access_automatically(tmp_path: Path):
    app = create_app(
        database_url=f"sqlite:///{tmp_path/'profile.db'}",
        storage_root=tmp_path/"storage",
        qr_secret="s",
        public_base_url="http://testserver",
        stt_url=None,
    )
    client = TestClient(app)

    admin = {
        "X-Tenant-ID": "t1",
        "X-User-ID": "admin",
        "X-Branch-ID": "b1",
        "X-Practitioner-ID": "admin-p",
    }
    doctor = {
        "X-Tenant-ID": "t1",
        "X-User-ID": "doctor-user",
        "X-Branch-ID": "b1",
        "X-Practitioner-ID": "doctor-1",
    }

    profile = client.post("/v1/admin/doctor-profiles", headers=admin, json={
        "practitioner_id": "doctor-1",
        "display_name": "Иванова Анна",
        "specialty_codes": ["GYNE"]
    })
    assert profile.status_code == 201, profile.text

    source = tmp_path/"gyne.docx"
    d = Document()
    d.add_paragraph("Жалобы: {{complaints}}")
    d.save(source)

    template = client.post(
        "/v1/admin/templates",
        headers=admin,
        files={"file": ("gyne.docx", source.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={
            "name": "Осмотр гинеколога",
            "fields_json": json.dumps([{"id":"complaints","type":"textarea"}]),
            "assignments_json": json.dumps([{"specialty_code":"GYNE","branch_id":"b1"}]),
        }
    ).json()["data"]
    client.post(f"/v1/admin/templates/{template['id']}/publish", headers=admin)

    eligible = client.get("/v1/doctor/templates", headers=doctor)
    assert eligible.status_code == 200, eligible.text
    assert [row["id"] for row in eligible.json()["data"]] == [template["id"]]

    created_document = client.post("/v1/doctor/documents", headers=doctor, json={
        "template_id": template["id"],
        "patient_id": "patient-1",
        "visit_type": None,
        "system_values": {}
    })
    assert created_document.status_code == 201, created_document.text

    wrong = dict(doctor)
    wrong["X-Practitioner-ID"] = "doctor-2"
    assert client.get("/v1/doctor/templates", headers=wrong).json()["data"] == []
